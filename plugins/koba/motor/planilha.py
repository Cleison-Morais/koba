"""Koba - monta a planilha de notas a partir dos dados coletados na plataforma.

Pastas (dentro da pasta de trabalho do Koba):
  dados/<curso>.json   -> resumo coletado pelo coletor.js (um por disciplina)
  estado.json          -> notas que o robô deu aos registros (rascunho / publicado)
  config.json          -> disciplinas, pesos e semanas que contam
  entrada/             -> solte aqui a planilha da Prova Paulista (.xlsx ou .csv)
Saída: a planilha definida em config.json ("planilha"), na mesma pasta.
Mantém o que o professor digitou nas colunas de notas manuais (Prova Paulista e a "outra nota").
"""
import csv, glob, json, os, re, shutil, sys, unicodedata
from datetime import datetime
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
_CFG = json.load(open(os.path.join(AQUI, "config.json"), encoding="utf-8")) if os.path.exists(os.path.join(AQUI, "config.json")) else {}
SAIDA = os.path.join(AQUI, _CFG.get("planilha", "NOTAS KOBA.xlsx"))
NOMES_OUTRA = {"Classroom + Soft skill", "Soft skill", "Outra nota", _CFG.get("nome_outra_nota", "Outra nota")} | set(_CFG.get("nome_classroom_por_aba", {}).values())
CORTE = str(_CFG.get("nota_corte", 5))


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def carregar(nome, padrao):
    p = os.path.join(AQUI, nome)
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else padrao


def semana_num(tarefa):
    m = re.match(r"S(\d+)", tarefa)
    return int(m.group(1)) if m else 0


def ler_manuais():
    """Lê Prova Paulista e Classroom já digitados na planilha atual."""
    manuais = {}
    if not os.path.exists(SAIDA):
        return manuais
    wb = openpyxl.load_workbook(SAIDA, data_only=False)
    for ws in wb.worksheets:
        cab = [c.value for c in ws[1]]
        if "Aluno" not in cab or "Prova Paulista" not in cab or not (NOMES_OUTRA & set(cab)):
            continue
        ia, ip = cab.index("Aluno"), cab.index("Prova Paulista")
        ic = next(i for i, c in enumerate(cab) if c in NOMES_OUTRA)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[ia]:
                manuais.setdefault(ws.title, {})[norm(row[ia])] = {"pp": row[ip], "cls": row[ic]}
    return manuais


def importar_prova_paulista(manuais, abas_alunos):
    """Importa notas de qualquer planilha solta em entrada/ e move o arquivo para entrada/importados."""
    pasta = os.path.join(AQUI, "entrada")
    os.makedirs(os.path.join(pasta, "importados"), exist_ok=True)
    msgs = []
    for arq in glob.glob(os.path.join(pasta, "*.xlsx")) + glob.glob(os.path.join(pasta, "*.csv")):
        linhas = []
        if arq.endswith(".csv"):
            with open(arq, encoding="utf-8-sig", errors="ignore") as f:
                amostra = f.read(2048); f.seek(0)
                sep = ";" if amostra.count(";") > amostra.count(",") else ","
                linhas = list(csv.reader(f, delimiter=sep))
        else:
            ws = openpyxl.load_workbook(arq, data_only=True).active
            linhas = [list(r) for r in ws.iter_rows(values_only=True)]
        # acha a linha de cabeçalho com uma coluna de nome e uma de nota
        ic = inota = None
        for i, lin in enumerate(linhas[:15]):
            cab = [norm(c) for c in lin]
            cn = [j for j, c in enumerate(cab) if c in ("nome", "aluno", "nome do aluno", "estudante", "nome do estudante")]
            cv = [j for j, c in enumerate(cab) if any(k in c for k in ("nota", "proficiencia", "acertos", "pontuacao", "resultado", "prova paulista"))]
            if cn and cv:
                ic, inota, ini = cn[0], cv[-1], i + 1
                break
        if ic is None:
            msgs.append(f"Não entendi as colunas de {os.path.basename(arq)} (preciso de uma coluna 'Nome' e uma 'Nota').")
            continue
        notas = {}
        for lin in linhas[ini:]:
            if len(lin) > max(ic, inota) and lin[ic]:
                try:
                    notas[norm(lin[ic])] = float(str(lin[inota]).replace(",", "."))
                except ValueError:
                    pass
        maior = max(notas.values()) if notas else 10
        fator = 10 / maior if maior > 10 else 1  # converte para 0-10 se vier em outra escala
        # a disciplina pode estar no nome do arquivo; senão vale para todas
        alvo = [a for a in abas_alunos if norm(a)[:10] in norm(os.path.basename(arq))] or list(abas_alunos)
        achados = 0
        for aba in alvo:
            for aluno in abas_alunos[aba]:
                if norm(aluno) in notas:
                    manuais.setdefault(aba, {}).setdefault(norm(aluno), {"pp": None, "cls": None})["pp"] = round(notas[norm(aluno)] * fator, 1)
                    achados += 1
        shutil.move(arq, os.path.join(pasta, "importados", datetime.now().strftime("%Y%m%d-%H%M_") + os.path.basename(arq)))
        msgs.append(f"Prova Paulista importada de {os.path.basename(arq)}: {achados} notas lançadas em {', '.join(alvo)}.")
    return msgs


def main():
    cfg = carregar("config.json", {})
    estado = carregar("estado.json", {})
    pesos_padrao = cfg.get("pesos", {"registro": 0.2, "quiz": 0.1, "prova_paulista": 0.3, "classroom": 0.4})
    sem_ini, sem_fim = cfg.get("semana_inicial", 0), cfg.get("semana_final", 99)

    dados = []
    for p in sorted(glob.glob(os.path.join(AQUI, "dados", "[0-9]*.json"))):
        d = json.load(open(p, encoding="utf-8"))
        if isinstance(d.get("alunos"), list):  # formato compacto do coletor
            # a plataforma às vezes esconde quem não enviou nada: completa com a lista de alunos já conhecida
            rp = os.path.join(AQUI, "dados", f"alunos_{d['curso']}.json")
            roster = json.load(open(rp, encoding="utf-8")) if os.path.exists(rp) else {}
            nt = len(d["tarefas"])
            for a in d["alunos"]:
                a[2] = [(-1 if v is None else v) for v in a[2]] + [-1] * (nt - len(a[2]))
                roster[str(a[0])] = a[1]
            vistos = {str(a[0]) for a in d["alunos"]}
            for uid, nome in roster.items():
                if uid not in vistos:
                    d["alunos"].append([int(uid), nome, [-1] * nt, [None] * len(d["semanas"])])
            json.dump(roster, open(rp, "w", encoding="utf-8"), ensure_ascii=False)
            tar = [{"tarefa": t[0], "cmid": t[1], "assignmentid": t[2], "contextid": t[3]} for t in d["tarefas"]]
            al = {}
            for uid, nome, regs, qs in d["alunos"]:
                al[str(uid)] = {"nome": nome,
                                "reg": {t["tarefa"]: ([0, None] if v == -1 else [1, None] if v == -2 else [1, v]) for t, v in zip(tar, regs)},
                                "quizsem": {s: q for s, q in zip(d["semanas"], qs)}}
            d["tarefas"], d["alunos"] = tar, al
        d["aba"] = cfg.get("abas", {}).get(str(d["curso"]), d["disciplina"].replace("Disciplina:", "").split("–")[0].strip())[:31]
        dados.append(d)

    manuais = ler_manuais()
    abas_alunos = {d["aba"]: [a["nome"] for a in d["alunos"].values()] for d in dados}
    avisos = importar_prova_paulista(manuais, abas_alunos)
    # cópia de segurança da Prova Paulista: se a planilha voltar sem as notas, recupera daqui
    bkp_path = os.path.join(AQUI, "dados", "prova_paulista_backup.json")
    bkp = carregar(os.path.join("dados", "prova_paulista_backup.json"), {}) if os.path.exists(bkp_path) else {}
    recuperadas = {}
    for aba, alunos in bkp.items():
        for nome, pp in alunos.items():
            m = manuais.setdefault(aba, {}).setdefault(nome, {"pp": None, "cls": None})
            if m.get("pp") in (None, "") and pp is not None:
                m["pp"] = pp
                recuperadas[aba] = recuperadas.get(aba, 0) + 1
    for aba, n in recuperadas.items():
        avisos.append(f"{n} notas da Prova Paulista de {aba} tinham sumido da planilha e foram recuperadas da cópia de segurança.")
    for aba, alunos in manuais.items():
        for nome, m in alunos.items():
            if isinstance(m.get("pp"), (int, float)):
                bkp.setdefault(aba, {})[nome] = m["pp"]
    with open(bkp_path, "w", encoding="utf-8") as f:
        json.dump(bkp, f, ensure_ascii=False, indent=1)

    wb = openpyxl.Workbook(); wb.remove(wb.active)
    azul, amarelo, vermelho, verde = (PatternFill("solid", fgColor=c) for c in ("1F4E78", "FFF2CC", "F8CBAD", "E2EFDA"))
    fina = Side(style="thin", color="BFBFBF"); borda = Border(top=fina, bottom=fina, left=fina, right=fina)
    resumo = []
    cobrar = {}

    # notas da atividade de recuperação: substituem Registro + Quiz quando forem maiores (célula fica amarela)
    rec_path = os.path.join(AQUI, "dados", "recuperacao.json")
    recup = {aba: {norm(n): v for n, v in notas.items()} for aba, notas in (json.load(open(rec_path, encoding="utf-8")).items() if os.path.exists(rec_path) else [])}
    usadas_rec = []

    for d in dados:
        pesos = {**pesos_padrao, **cfg.get("pesos_por_aba", {}).get(d["aba"], {})}
        nome_cls = cfg.get("nome_classroom_por_aba", {}).get(d["aba"], cfg.get("nome_outra_nota", "Outra nota"))
        ws = wb.create_sheet(d["aba"])
        tarefas = sorted([t for t in d["tarefas"] if sem_ini <= semana_num(t["tarefa"]) <= sem_fim], key=lambda t: semana_num(t["tarefa"]))
        semanas_quiz = sorted({semana_num(t["tarefa"]) for t in tarefas})
        cab = ["Aluno"] + [f"Reg {t['tarefa']}" for t in tarefas] + [f"Quiz S{s}" for s in semanas_quiz] + [
            "Média registros", "Média quiz", "Registro + Quiz", "Prova Paulista", nome_cls, "MÉDIA FINAL", "Não entregou"]
        for i, h in enumerate(cab, 1):
            c = ws.cell(1, i, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = azul
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        nR, nQ = len(tarefas), len(semanas_quiz)
        cR0, cQ0 = 2, 2 + nR
        cMR = 2 + nR + nQ
        L = get_column_letter
        alunos = sorted(d["alunos"].items(), key=lambda kv: norm(kv[1]["nome"]))
        fizeram = {t["tarefa"]: [] for t in tarefas}; nao = {t["tarefa"]: [] for t in tarefas}
        for r, (uid, a) in enumerate(alunos, 2):
            ws.cell(r, 1, a["nome"])
            falta = []
            for j, t in enumerate(tarefas):
                env, nota = a["reg"].get(t["tarefa"], [0, None])
                c = ws.cell(r, cR0 + j)
                chave = f"{t['cmid']}:{uid}"
                if nota is not None:
                    c.value = nota
                elif env and chave in estado:
                    c.value = estado[chave]["nota"]; c.fill = amarelo  # nota do robô aguardando/depois de aprovação
                elif env:
                    c.value = None; c.fill = amarelo
                else:
                    c.value = 0; c.fill = vermelho; falta.append(t["tarefa"])
                (fizeram if env else nao)[t["tarefa"]].append(a["nome"])
            for j, s in enumerate(semanas_quiz):
                ws.cell(r, cQ0 + j, a.get("quizsem", {}).get(s))
            m = manuais.get(d["aba"], {}).get(norm(a["nome"]), {})
            rR = f"{L(cR0)}{r}:{L(cR0+nR-1)}{r}"; rQ = f"{L(cQ0)}{r}:{L(cQ0+nQ-1)}{r}"
            ws.cell(r, cMR, f'=IFERROR(ROUND(AVERAGE({rR}),1),0)' if nR else 0)
            ws.cell(r, cMR + 1, f'=IFERROR(ROUND(AVERAGE({rQ}),1),0)' if nQ else 0)
            wr, wq = pesos['registro'], pesos['quiz']
            wrq, wpp, wcl = round(wr + wq, 4), pesos["prova_paulista"], pesos["classroom"]
            # média de registro e quiz numa nota só (0 a 10), com o registro pesando mais
            ws.cell(r, cMR + 2, f"=ROUND(({L(cMR)}{r}*{wr}+{L(cMR+1)}{r}*{wq})/{wrq},1)")
            nr = recup.get(d["aba"], {}).get(norm(a["nome"]))
            if nr is not None:
                vr = [ws.cell(r, cR0 + j).value for j in range(nR)]; vr = [x for x in vr if isinstance(x, (int, float))]
                vq = [ws.cell(r, cQ0 + j).value for j in range(nQ)]; vq = [x for x in vq if isinstance(x, (int, float))]
                mr = round(sum(vr) / len(vr), 1) if vr else 0; mq = round(sum(vq) / len(vq), 1) if vq else 0
                antes = round((mr * wr + mq * wq) / wrq, 1)
                if nr > antes:
                    c = ws.cell(r, cMR + 2, nr); c.fill = amarelo
                    c.comment = Comment(f"Nota da recuperação: {nr}. Antes era {antes} (registro + quiz).", "Robô de notas")
                usadas_rec.append({"aba": d["aba"], "nome": a["nome"], "recuperacao": nr, "antes": antes, "substituiu": nr > antes})
            ws.cell(r, cMR + 3, m.get("pp"))
            ws.cell(r, cMR + 4, m.get("cls"))
            # média final já com os pesos (registro+quiz, prova paulista, classroom); vazio conta 0
            ws.cell(r, cMR + 5, f'=ROUND({L(cMR+2)}{r}*{wrq}+N({L(cMR+3)}{r})*{wpp}+N({L(cMR+4)}{r})*{wcl},1)')
            ws.cell(r, cMR + 6, ", ".join(falta))
            if falta:
                cobrar.setdefault(a["nome"], []).append(f"{d['aba']}: {', '.join(falta)}")
            for col in range(1, len(cab) + 1):
                c = ws.cell(r, col); c.border = borda
                if col > 1 and col < len(cab):
                    c.alignment = Alignment(horizontal="center"); c.number_format = "0.0"
                if col in (cMR + 3, cMR + 4) and c.fill.fgColor.rgb in (None, "00000000"):
                    c.fill = verde
            ws.cell(r, cMR + 5).font = Font(bold=True)
        ws.column_dimensions["A"].width = 38
        for col in range(2, len(cab) + 1):
            ws.column_dimensions[L(col)].width = 11
        ws.column_dimensions[L(len(cab))].width = 30
        ws.row_dimensions[1].height = 45; ws.freeze_panes = "B2"
        for col in (cMR + 2, cMR + 3, cMR + 4, cMR + 5):
            ws.column_dimensions[L(col)].width = 15
        # notas de cada registro e quiz ficam agrupadas e recolhidas (botão + em cima para abrir)
        for col in range(cR0, cMR + 2):
            ws.column_dimensions[L(col)].outlineLevel = 1
            ws.column_dimensions[L(col)].hidden = True
        ws.sheet_properties.outlinePr.summaryRight = True
        rm = len(alunos) + 2
        ws.cell(rm, 1, "MÉDIA DA TURMA").font = Font(bold=True)
        for col in range(2, len(cab)):
            L2 = get_column_letter(col)
            c = ws.cell(rm, col, f'=IFERROR(ROUND(AVERAGE({L2}2:{L2}{rm-1}),1),"")')
            c.font = Font(bold=True); c.number_format = "0.0"; c.alignment = Alignment(horizontal="center")
            c.fill = PatternFill("solid", fgColor="D9E1F2")
        ws.cell(rm, 1).fill = PatternFill("solid", fgColor="D9E1F2")
        # cor da MÉDIA FINAL: igual à nota de corte = amarelo, abaixo = vermelho, acima = azul
        from openpyxl.formatting.rule import CellIsRule
        faixa = f"{L(cMR+5)}2:{L(cMR+5)}{rm-1}"
        ws.conditional_formatting.add(faixa, CellIsRule(operator="equal", formula=[CORTE], fill=PatternFill("solid", fgColor="FFE34D", bgColor="FFE34D"), font=Font(bold=True, color="000000")))
        ws.conditional_formatting.add(faixa, CellIsRule(operator="lessThan", formula=[CORTE], fill=PatternFill("solid", fgColor="F28B82", bgColor="F28B82"), font=Font(bold=True, color="7F0000")))
        ws.conditional_formatting.add(faixa, CellIsRule(operator="greaterThan", formula=[CORTE], fill=PatternFill("solid", fgColor="8AB4F8", bgColor="8AB4F8"), font=Font(bold=True, color="0B2A6B")))
        n = len(alunos) + 4
        legenda = ["Vermelho = não entregou (conta 0).  Amarelo = nota dada pelo robô aguardando publicação; em Registro + Quiz, amarelo = nota substituída pela recuperação.",
                   "Verde = colunas que você preenche (0 a 10). A Prova Paulista também entra sozinha se você soltar a planilha na pasta entrada/ do Koba.",
                   f"Atualizado em {datetime.now():%d/%m/%Y %H:%M}. Dados da plataforma coletados em {d['coletadoEm'][:16].replace('T', ' ')}."]
        for i, t in enumerate(legenda):
            ws.cell(n + i, 1, t).font = Font(italic=True, color="808080")
        if tarefas:
            ult = tarefas[-1]["tarefa"]
            resumo.append((d["aba"], ult, fizeram[ult], nao[ult]))

    wr = wb.create_sheet("Resumo", 0)
    wr["A1"] = f"Resumo de entregas - {datetime.now():%d/%m/%Y %H:%M}"; wr["A1"].font = Font(bold=True, size=14)
    lin = 3
    for aba, ult, sim, nao_ in resumo:
        wr.cell(lin, 1, f"{aba} - {ult}").font = Font(bold=True); lin += 1
        wr.cell(lin, 1, f"Entregaram ({len(sim)})").font = Font(color="2E7D32", bold=True)
        wr.cell(lin, 2, f"Não entregaram ({len(nao_)})").font = Font(color="C62828", bold=True); lin += 1
        for i in range(max(len(sim), len(nao_))):
            wr.cell(lin + i, 1, sim[i] if i < len(sim) else None); wr.cell(lin + i, 2, nao_[i] if i < len(nao_) else None)
        lin += max(len(sim), len(nao_)) + 2
    wr.cell(lin, 1, "PARA COBRAR - atividades de registro não entregues").font = Font(bold=True, size=13, color="C62828"); lin += 1
    for nome in sorted(cobrar, key=norm):
        wr.cell(lin, 1, nome); wr.cell(lin, 2, " | ".join(cobrar[nome])); wr.cell(lin, 2).alignment = Alignment(wrap_text=True); lin += 1
    lin += 1
    for i, a in enumerate(avisos):
        wr.cell(lin + i, 1, a).font = Font(italic=True)
    wr.column_dimensions["A"].width = 45; wr.column_dimensions["B"].width = 90

    try:
        wb.save(SAIDA)
    except PermissionError:
        print("ERRO: a planilha está aberta no Excel. Feche e rode de novo."); sys.exit(2)
    print(json.dumps({"ok": True, "saida": SAIDA, "avisos": avisos, "recuperacao": usadas_rec, "cobrar": cobrar,
                      "resumo": [{"disciplina": a, "tarefa": u, "entregaram": len(s), "nao_entregaram": len(n), "lista_nao_entregaram": n} for a, u, s, n in resumo]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
